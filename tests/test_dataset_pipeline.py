"""Tests for local dataset persistence, exports, and the scraper CLI."""

import json
from argparse import Namespace
from pathlib import Path

from openpyxl import load_workbook

from app.integrations.dataset_export import export_json, export_xlsx
from app.repositories.dataset import LocalDatasetRepository
from app.schemas.scraper import ScrapedQAPair, ThreadContextMessage
from app.services.dataset import DatasetService
from scripts.scrape_kb import run


def record(score: float = 0.5) -> ScrapedQAPair:
    """Build a valid synthetic Q&A record."""
    return ScrapedQAPair(
        question="How do I sign in?",
        answer="Use the account portal and follow the sign-in steps.",
        confidence_score=score,
        is_trusted_answer=score >= 0.8,
        source_channel="support-general",
        thread_ts="1700000000.000000",
        asker_name="Asker",
        answerer_name="Answerer",
        date="2023-11-14",
        full_thread_context=[
            ThreadContextMessage(
                author="Answerer", text="Use the portal.", is_mentor=True
            )
        ],
    )


def test_repository_round_trip_and_replace(tmp_path: Path):
    """Persist validated records and replace the previous snapshot."""
    repository = LocalDatasetRepository(tmp_path / "dataset.json")
    repository.save([record()])

    assert repository.load()[0].answerer_name == "Answerer"
    repository.save([record(score=1.0)])
    assert repository.load()[0].confidence_score == 1.0


def test_dataset_service_sorts_and_exports_json_and_xlsx(tmp_path: Path):
    """Publish strongest records to machine and human review formats."""
    repository = LocalDatasetRepository(tmp_path / "snapshot.json")
    service = DatasetService(repository)
    json_path = tmp_path / "output" / "dataset.json"
    xlsx_path = tmp_path / "output" / "dataset.xlsx"

    published = service.publish([record(0.2), record(0.9)], json_path, xlsx_path)

    assert [item.confidence_score for item in published] == [0.9, 0.2]
    assert json.loads(json_path.read_text())[0]["full_thread_context"]
    workbook = load_workbook(xlsx_path)
    assert workbook.active.max_row == 3
    assert workbook.active.max_column == 9


def test_export_helpers_create_parent_directories(tmp_path: Path):
    """Create ignored output directories when exporting directly."""
    json_path = tmp_path / "nested" / "records.json"
    xlsx_path = tmp_path / "nested" / "records.xlsx"

    export_json([record()], json_path)
    export_xlsx([record()], xlsx_path)

    assert json_path.exists()
    assert xlsx_path.exists()


class DryRunReader:
    """Minimal reader proving CLI discovery does not fetch message history."""

    def fetch_bot_user_id(self):
        """Return a synthetic bot identity."""
        return "U_BOT"

    def list_channels(self):
        """Return one matching and one unrelated channel."""
        return [
            {"name": "support-general", "id": "C1", "is_member": True},
            {"name": "random", "id": "C2", "is_member": True},
        ]


def test_cli_dry_run_discovers_without_fetching_messages(capsys, tmp_path: Path):
    """List targets without requiring Slack history access or output files."""
    args = Namespace(
        dry_run=True,
        channels=None,
        exclude_channels=[],
        patterns="support-",
        mentors_channel="experts",
        output_dir=tmp_path,
    )

    assert run(args, DryRunReader()) == 0
    assert "support-general" in capsys.readouterr().out
    assert not list(tmp_path.iterdir())
